"""Tests de asistencia/export.py · spec 007."""

from datetime import timezone
from io import BytesIO

import pytest
from asistencia.export import report_heading, safe_filename, to_xlsx
from openpyxl import load_workbook


def rows(*pairs):
    return [{"name": name, "checked_at": at} for name, at in pairs]


def workbook_cells(content: bytes):
    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = book.active
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def test_report_heading():
    assert report_heading("Clase 20/10 – SDD", "2026-10-24") == (
        "Asistencia Clase 20/10 – SDD - 2026-10-24"
    )


def test_safe_filename_sanitizes_slashes():
    heading = report_heading("Clase 20/10 – SDD", "2026-10-24")
    assert safe_filename(heading) == "Asistencia Clase 20-10 – SDD - 2026-10-24.xlsx"


def test_safe_filename_no_session_id_suffix():
    name = safe_filename(report_heading("Mañana", "2026-10-24"))
    assert "jUGLyk0n" not in name
    assert name.endswith(".xlsx")


def test_title_and_total_use_four_merged_columns():
    heading = report_heading("Centrado", "2026-10-24")
    content = to_xlsx(heading, [], tz=timezone.utc)
    book = load_workbook(BytesIO(content), read_only=False)
    ws = book.active
    merged = {str(area) for area in ws.merged_cells.ranges}
    assert merged == {"A1:D1", "A2:D2"}
    assert ws["A1"].alignment.horizontal == "center"
    assert ws["A2"].alignment.horizontal == "left"


def test_empty_session_report():
    heading = report_heading("Vacía", "2026-10-24")
    cells = workbook_cells(to_xlsx(heading, [], tz=timezone.utc))
    assert cells[0][0] == heading
    assert cells[1][0] == "Total de asistentes: 0"
    assert cells[3][:4] == ["#", "Nombre", "Fecha", "Hora"]
    assert len(cells) == 4


def test_rows_sorted_and_split_datetime():
    heading = report_heading("Orden", "2026-10-20")
    content = to_xlsx(
        heading,
        rows(
            ("Luis", "2026-10-20T15:05:00+00:00"),
            ("Ana García", "2026-10-20T15:01:30+00:00"),
        ),
        tz=timezone.utc,
    )
    cells = workbook_cells(content)
    assert cells[1][0] == "Total de asistentes: 2"
    assert [cells[4][1], cells[5][1]] == ["Ana García", "Luis"]
    assert cells[4][2:] == ["2026-10-20", "15:01:30"]
    assert cells[5][2:] == ["2026-10-20", "15:05:00"]


def test_special_characters_in_name_preserved():
    heading = report_heading("X", "2026-10-20")
    content = to_xlsx(
        heading,
        rows(('Ana; "la jefa"', "2026-10-20T15:00:00+00:00")),
        tz=timezone.utc,
    )
    cells = workbook_cells(content)
    assert cells[4][1] == 'Ana; "la jefa"'


@pytest.mark.parametrize("dangerous", ["=SUMA(A1)", "+34", "-1", "@cmd"])
def test_formula_like_names_stored_as_text(dangerous):
    heading = report_heading("X", "2026-10-20")
    content = to_xlsx(
        heading,
        rows((dangerous, "2026-10-20T15:00:00+00:00")),
        tz=timezone.utc,
    )
    book = load_workbook(BytesIO(content), read_only=False)
    cell = book.active["B5"]
    assert cell.value == dangerous
    assert cell.number_format == "@"
    assert cell.quotePrefix is True
