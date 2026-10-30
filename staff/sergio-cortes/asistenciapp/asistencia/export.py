"""Exportación de la lista de asistencia a informe `.xlsx` (spec 007).

Se llama export.py y no csv.py para no tapar el módulo `csv` estándar.
"""

import re
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

TABLE_COLUMNS = 4  # #, Nombre, Fecha, Hora (A–D)

FORMULA_PREFIXES = ("=", "+", "-", "@")
_FILENAME_BAD = re.compile(r'[\\/:*?"<>|]')
_MAX_FILENAME_LEN = 200


def report_heading(title: str, opened_date: str) -> str:
    """Texto del encabezado del informe y base del nombre de archivo.

    :param title: título de la sesión.
    :param opened_date: fecha de apertura ``AAAA-MM-DD``.
    :return: ``Asistencia {title} - {opened_date}``.
    :rtype: str
    """
    return f"Asistencia {title} - {opened_date}"


def safe_filename(heading: str) -> str:
    """Nombre de archivo seguro para Mac/Windows con extensión ``.xlsx``.

    :param heading: texto del encabezado (sin extensión).
    :return: nombre sanitizado terminado en ``.xlsx``.
    :rtype: str
    """
    stem = _FILENAME_BAD.sub("-", heading.strip())
    if not stem:
        stem = "Asistencia"
    if len(stem) > _MAX_FILENAME_LEN:
        stem = stem[:_MAX_FILENAME_LEN].rstrip("- ")
    return f"{stem}.xlsx"


def _timezone(tz):
    if tz is not None:
        return tz
    return datetime.now().astimezone().tzinfo


def _split_checked_at(iso: str, tz) -> tuple[str, str]:
    moment = datetime.fromisoformat(iso).astimezone(tz)
    return moment.strftime("%Y-%m-%d"), moment.strftime("%H:%M:%S")


def _merge_row(ws, row: int, value: str, *, bold: bool, horizontal: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=TABLE_COLUMNS)
    cell = ws.cell(row=row, column=1, value=value)
    cell.font = Font(bold=bold)
    cell.alignment = Alignment(horizontal=horizontal, vertical="center", wrap_text=True)


def _write_name_cell(cell, name: str) -> None:
    cell.value = name
    if name.startswith(FORMULA_PREFIXES):
        cell.number_format = "@"
        cell.quotePrefix = True


def to_xlsx(heading: str, rows, tz=None) -> bytes:
    """Genera un informe xlsx con título, total y tabla de asistentes.

    :param heading: encabezado visible (fila 1).
    :param rows: dicts con ``name`` y ``checked_at`` (ISO UTC).
    :param tz: zona horaria para fechas de fila; ``None`` → local del servidor.
    :return: contenido binario del libro.
    :rtype: bytes
    """
    zone = _timezone(tz)
    wb = Workbook()
    ws = wb.active
    ws.title = "Asistencia"

    _merge_row(ws, 1, heading, bold=True, horizontal="center")
    _merge_row(
        ws,
        2,
        f"Total de asistentes: {len(rows)}",
        bold=False,
        horizontal="left",
    )

    for col, label in enumerate(["#", "Nombre", "Fecha", "Hora"], start=1):
        cell = ws.cell(row=4, column=col, value=label)
        cell.font = Font(bold=True)

    ordered = sorted(rows, key=lambda row: row["checked_at"])
    for index, row in enumerate(ordered, start=1):
        excel_row = 4 + index
        date_part, time_part = _split_checked_at(row["checked_at"], zone)
        ws.cell(row=excel_row, column=1, value=index)
        _write_name_cell(ws.cell(row=excel_row, column=2), row["name"])
        ws.cell(row=excel_row, column=3, value=date_part)
        ws.cell(row=excel_row, column=4, value=time_part)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
